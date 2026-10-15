"""`apartmentscout export --xlsx FILE | --csv FILE`: the search as a scored workbook or a flat CSV.

The workbook is a working tracker, not a static dump:

- **Weights**: one editable cell per scoring dimension, the neutral value for unknown dimensions and the target
  monthly cost. Changing a cell re-scores and re-ranks the Listings sheet (live SUMPRODUCT and COUNTIFS formulas).
- **Listings**: every active listing with its seven 0-100 dimension scores stored as values, plus live Score, Rank,
  $/sq ft and Vs target formulas; frozen headers, filters, number formats and colour scales.
- **Outreach** (questions per building, with blank tracker columns), **To verify** (missing facts), **Upfront
  charges**, **Rejected or held**, **Gone fast** (only when listing history exists) and **Sources**.

Every text cell goes through the formula-injection guard (see apartmentscout/handoff/safety.py), in the CSV too.
openpyxl is optional: ``pip install -e ".[xlsx]"``.
"""
from __future__ import annotations

import copy
import csv
import datetime as dt
import json
import sys
from pathlib import Path
from typing import Any, Callable, Iterable

from .. import renterfacts, tours as tours_mod
from ..handoff.facts import focus_listings, is_active, missing_facts, parse_date, questions_for
from ..handoff.safety import safe_cell

try:
    from ..config import DIMENSION_LABELS as DIM_LABELS
except ImportError:  # older scorer without the shared label table
    DIM_LABELS = {"cost": "Cost", "space": "Space", "commute": "Commute", "amenities": "Amenities",
                  "light_view": "Light & view", "building": "Building", "quiet": "Quiet"}
DIM_HELP = {
    "cost": "Monthly cost (net-effective rent + fees) against the budget tiers and the hard cap",
    "space": "Square footage against sqft_floor, sqft_good and sqft_great; halved below min_beds",
    "commute": "Door-to-door minutes against each destination's max_minutes, weighted by destination",
    "amenities": "Must-haves (x3), strong preferences (x2) and nice-to-haves (x1); unknown counts half",
    "light_view": "Direct sun hours from the 3D site model, else exposure, floor and listing wording",
    "building": "Resident review ratings and public-record risk",
    "quiet": "311 noise around the building, else review wording",
    "move_date": "Availability against the move-in window",
    "renewal_risk": "How far the second-year rent could jump once a concession ends",
    "transit_access": "Distinct lines within a short walk, and one-seat rides to the destinations",
}
URL_SCHEMES = ("http://", "https://")


class Formula(str):
    """A string that is written as a formula (bypasses the injection guard). Only this module creates them."""


# --------------------------------------------------------------------------- command
def register(add: Callable[..., Any]) -> None:
    s = add("export", cmd_export, "export the search as a scored workbook (--xlsx) or a CSV (--csv)")
    s.add_argument("--xlsx", metavar="FILE", help="write an Excel workbook with live weights, scores and ranks")
    s.add_argument("--csv", metavar="FILE", help="write every listing as one CSV row (formula-injection safe)")
    s.add_argument("--top", type=int, default=25,
                   help="ranked listings to add to the Outreach sheet besides the shortlist (default 25)")


def cmd_export(args: Any) -> None:
    from ..cli import log, ws_from
    from ..listings import load
    if not args.xlsx and not args.csv:
        sys.exit("Nothing to export: pass --xlsx FILE and/or --csv FILE.")
    ws = ws_from(args)
    crit = ws.criteria()
    Ls = scored(load(ws.listings_path) if ws.listings_path.exists() else [], crit, tours_mod.load_all(ws.root), renterfacts.load(ws.root))
    if args.csv:
        n = write_csv(Ls, crit, Path(args.csv))
        log(f"CSV: {n} listings -> {Path(args.csv).resolve()}")
    if args.xlsx:
        sheets = write_xlsx(Ls, crit, Path(args.xlsx), history=load_gone_fast(ws, Ls), top=args.top)
        log(f"Workbook: {', '.join(f'{k} ({v})' for k, v in sheets.items())} -> {Path(args.xlsx).resolve()}")


def scored(listings: list[dict], crit: dict, tours: dict[str, dict] | None = None, facts: dict[str, dict] | None = None) -> list[dict]:
    """Score a copy with the current criteria, so the export never depends on a stale `score` and never writes."""
    from ..score import score_all
    Ls = copy.deepcopy(listings)
    score_all(Ls, crit, tours=tours, facts=facts)
    return Ls


def ordered(listings: Iterable[dict]) -> list[dict]:
    """Ranked first (by rank), then other active listings by score, then rejected, then gone."""
    def key(L: dict) -> tuple:
        s = L.get("score") or {}
        group = 0 if s.get("rank") else 1 if is_active(L) else 2 if L.get("status") == "rejected" else 3
        return (group, s.get("rank") or 0, -(s.get("overall") or -1), L["id"])
    return sorted(listings, key=key)


# --------------------------------------------------------------------------- columns
def _dests(crit: dict) -> list[str]:
    return [d.get("name") for d in (crit.get("commute") or {}).get("destinations") or [] if d.get("name")]


def _date_or_text(v: Any) -> Any:
    d = parse_date(v)
    if d and len(str(v).strip()) <= 10:
        return d
    return v if v not in (None, "") else None


def _url(L: dict) -> str | None:
    u = (L.get("links") or {}).get("listing")
    return u if isinstance(u, str) and u.startswith(URL_SCHEMES) else None


def listing_columns(crit: dict) -> list[tuple[str, str, Callable[[dict], Any]]]:
    """(header, kind, value getter). Kinds drive number formats and widths; `*f` kinds become formulas in xlsx."""
    from ..listings import net_effective

    def price(k: str) -> Callable[[dict], Any]:
        return lambda L: (L.get("price") or {}).get(k)

    def part(k: str) -> Callable[[dict], Any]:
        return lambda L: ((L.get("score") or {}).get("parts") or {}).get(k)

    def mins(name: str) -> Callable[[dict], Any]:
        return lambda L: ((L.get("commute") or {}).get(name) or {}).get("minutes")

    target = _target(crit)
    cols: list[tuple[str, str, Callable[[dict], Any]]] = [
        ("Rank", "rankf", lambda L: (L.get("score") or {}).get("rank")),
        ("Score", "scoref", lambda L: (L.get("score") or {}).get("overall")),
        ("Tier", "text", lambda L: (L.get("score") or {}).get("budget_tier")),
        ("Building", "text", lambda L: L.get("building")),
        ("Unit", "text", lambda L: L.get("unit")),
        ("Neighborhood", "text", lambda L: L.get("neighborhood")),
        ("Status", "text", lambda L: L.get("status")),
        ("Eligible", "text", lambda L: "yes" if (L.get("score") or {}).get("eligible") else "no"),
        ("Monthly cost", "money", lambda L: (L.get("score") or {}).get("monthly_cost")),
        ("Vs target", "vstargetf", lambda L: (round(L["score"]["monthly_cost"] - target, 2)
                                             if target and (L.get("score") or {}).get("monthly_cost") else None)),
        ("Gross rent", "money", price("gross")),
        ("Net effective", "money", lambda L: net_effective(L.get("price") or {})),
        ("Fees / mo", "money", price("fees_monthly")),
        ("One-time fees", "money", price("fees_one_time")),
        ("Broker fee", "money", price("broker_fee")),
        ("Lease (mo)", "int", price("lease_months")),
        ("Free (mo)", "num1", price("free_months")),
        ("Beds", "num1", lambda L: L.get("beds")),
        ("Baths", "num1", lambda L: L.get("baths")),
        ("Sq ft", "sqft", lambda L: L.get("sqft")),
        ("$ / sq ft", "psff", lambda L: (round(L["score"]["monthly_cost"] / L["sqft"], 2)
                                         if L.get("sqft") and (L.get("score") or {}).get("monthly_cost") else None)),
        ("Floor", "int", lambda L: L.get("floor")),
        ("Exposure", "text", lambda L: " ".join(L.get("exposure") or []) or None),
    ]
    cols += [(f"{d} (min)", "int", mins(d)) for d in _dests(crit)]
    from ..score import DIMENSIONS
    cols += [(f"{DIM_LABELS.get(k, k)} score", "dim", part(k)) for k in DIMENSIONS]
    cols += [
        ("Confidence", "pct", lambda L: ((L.get("score") or {}).get("confidence") or 0) / 100),
        ("Range low", "num1", lambda L: ((L.get("score") or {}).get("range") or [None, None])[0]),
        ("Range high", "num1", lambda L: ((L.get("score") or {}).get("range") or [None, None])[1]),
        ("Gates", "text", lambda L: "; ".join((L.get("score") or {}).get("gates") or []) or None),
        ("Available", "date", lambda L: _date_or_text(L.get("available"))),
        ("Checked", "date", lambda L: _date_or_text(L.get("checked"))),
        ("Listing", "url", _url),
        ("ID", "text", lambda L: L.get("id")),
        ("Notes", "longtext", lambda L: L.get("notes")),
    ]
    return cols


def _target(crit: dict) -> float | None:
    tiers = (crit.get("budget") or {}).get("tiers") or []
    return tiers[0].get("max") if tiers else (crit.get("budget") or {}).get("hard_max")


# --------------------------------------------------------------------------- csv
def write_csv(listings: list[dict], crit: dict, path: Path) -> int:
    """Every listing, one row each, values only (Rank and Score from score.py). Text is injection-guarded."""
    cols = listing_columns(crit)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = ordered(listings)
    with path.open("w", newline="", encoding="utf-8-sig") as f:   # BOM so Excel reads UTF-8
        w = csv.writer(f)
        w.writerow([h for h, _, _ in cols])
        for L in rows:
            out = []
            for _, kind, get in cols:
                v = get(L)
                if isinstance(v, dt.date):
                    v = v.isoformat()
                elif kind == "pct" and v is not None:
                    v = round(v * 100)
                out.append(safe_cell(v) if v is not None else "")
            w.writerow(out)
    return len(rows)


# --------------------------------------------------------------------------- history
def load_gone_fast(ws: Any, listings: list[dict]) -> list[dict]:
    """Units that went off the market, with days listed when known.

    Uses the search's history (snapshots and claims under <search>/history/, via apartmentscout/quality/history.py)
    or a history.json in the search;
    otherwise listings marked gone that carry both a listed and a gone date. Empty when no history exists."""
    data: dict | None = None
    for p in (ws.root / "history.json", ws.root / "history" / "history.json"):
        if p.exists():
            try:
                data = json.loads(p.read_text(encoding="utf-8"))
                break
            except (OSError, ValueError):
                pass
    if data is None:
        try:
            from ..quality.history import history_doc  # the research pipeline's history (read-only call)
            data = history_doc(ws.root, listings, ws.criteria())
        except Exception:  # an optional history must never break the export
            data = None
    out: list[dict] = []
    if data and data.get("gone_fast"):
        out = [dict(g) for g in data["gone_fast"] if isinstance(g, dict) and g.get("id")]
    else:
        for L in listings:
            if L.get("status") != "gone":
                continue
            listed = parse_date(L.get("listed") or L.get("listed_at") or L.get("first_seen"))
            gone = parse_date(L.get("gone") or L.get("gone_date") or L.get("off_market"))
            if listed and gone:
                out.append({"id": L["id"], "listed": listed.isoformat(), "gone": gone.isoformat(),
                            "days": (gone - listed).days, "last_price": (L.get("price") or {}).get("gross")})
    return sorted(out, key=lambda g: (g.get("days") is None, g.get("days") or 0, g["id"]))


# --------------------------------------------------------------------------- xlsx
def _openpyxl() -> Any:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        sys.exit("The xlsx export needs openpyxl, an optional dependency. Install it with\n"
                 '    pip install -e ".[xlsx]"      (or: pip install openpyxl)\n'
                 "or export a CSV instead with --csv FILE.")
    import openpyxl
    return openpyxl


HEADER_ROW, FIRST_ROW = 4, 5
FORMATS = {"money": '"$"#,##0', "vstargetf": '+"$"#,##0;-"$"#,##0;"$"0', "psff": '"$"#,##0.00',
           "sqft": "#,##0", "int": "0", "num1": "0.0", "dim": "0.0", "scoref": "0.0", "rankf": "0",
           "pct": "0%", "date": "yyyy-mm-dd"}
WIDTHS = {"rankf": 6, "scoref": 7, "money": 11, "vstargetf": 10, "psff": 9, "sqft": 8, "int": 7, "num1": 7,
          "dim": 8, "pct": 9, "date": 11, "url": 18, "text": 16, "longtext": 48}
NAVY, INK, MUTED, EDIT = "1F2933", "FFFFFF", "6B7280", "FFF2CC"


class Book:
    """Small styling layer over openpyxl so every sheet looks the same."""

    def __init__(self) -> None:
        self.xl = _openpyxl()
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        self.wb = self.xl.Workbook()
        self.wb.remove(self.wb.active)
        self.Font, self.Fill, self.Align = Font, PatternFill, Alignment
        self.head_font = Font(bold=True, color=INK)
        self.head_fill = PatternFill("solid", fgColor=NAVY)
        self.edit_fill = PatternFill("solid", fgColor=EDIT)
        self.rule = Border(bottom=Side(style="thin", color="D1D5DB"))
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        self.illegal = ILLEGAL_CHARACTERS_RE

    def sheet(self, title: str, heading: str, caveat: str, headers: list[str], kinds: list[str]) -> Any:
        sh = self.wb.create_sheet(title)
        sh["A1"] = heading
        sh["A1"].font = self.Font(bold=True, size=14)
        sh["A2"] = caveat
        sh["A2"].font = self.Font(italic=True, color=MUTED)
        for i, (h, k) in enumerate(zip(headers, kinds, strict=False), 1):
            c = sh.cell(HEADER_ROW, i, h)
            c.font, c.fill = self.head_font, self.head_fill
            c.alignment = self.Align(wrap_text=True, vertical="center")
            sh.column_dimensions[self.xl.utils.get_column_letter(i)].width = WIDTHS.get(k, 14)
        sh.row_dimensions[HEADER_ROW].height = 32
        return sh

    def put(self, sh: Any, row: int, col: int, value: Any, kind: str = "text") -> Any:
        if isinstance(value, Formula):
            c = sh.cell(row, col)
            c.value = str(value)
        else:
            if isinstance(value, str):
                value = self.illegal.sub("", safe_cell(value))
            if kind == "url" and isinstance(value, str) and value.startswith(URL_SCHEMES):
                c = sh.cell(row, col, self.illegal.sub("", safe_cell(_host(value))))
                c.hyperlink = value
                c.style = "Hyperlink"
                return c
            c = sh.cell(row, col, value)
        if kind in FORMATS:
            c.number_format = FORMATS[kind]
        if kind == "longtext":
            c.alignment = self.Align(wrap_text=True, vertical="top")
        return c

    def finish(self, sh: Any, ncols: int, nrows: int, freeze: str = "A5") -> None:
        last = self.xl.utils.get_column_letter(max(ncols, 1))
        sh.auto_filter.ref = f"A{HEADER_ROW}:{last}{max(HEADER_ROW + nrows, HEADER_ROW + 1)}"
        sh.freeze_panes = freeze
        sh.sheet_view.zoomScale = 110


def _host(url: str) -> str:
    from urllib.parse import urlparse
    h = urlparse(url).netloc.lower()
    return (h[4:] if h.startswith("www.") else h) or "link"


def _col(xl: Any, i: int) -> str:
    return xl.utils.get_column_letter(i)


def _tier_colors(crit: dict) -> list[tuple[str, str]]:
    names = [t.get("name") for t in (crit.get("budget") or {}).get("tiers") or [] if t.get("name")]
    ramp = ["C6EFCE", "E2EFDA", "FFEB9C", "FCD5B4"]
    out = [(n, ramp[min(i * (len(ramp) - 1) // max(len(names) - 1, 1), len(ramp) - 1)]) for i, n in enumerate(names)]
    if names:
        out.append((f"Above {names[-1]}", "FCD5B4"))
    out.append(("Over budget", "FFC7CE"))
    return out


def write_xlsx(listings: list[dict], crit: dict, path: Path, history: list[dict] | None = None,
               top: int = 25) -> dict[str, int]:
    """Write the workbook. Returns {sheet title: data rows}."""
    B = Book()
    xl = B.xl
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from ..score import DIMENSIONS

    today = dt.date.today()
    rows = ordered(listings)
    active = [L for L in rows if is_active(L)]
    by_id = {L["id"]: L for L in listings}
    counts: dict[str, int] = {}
    search = (crit.get("search") or {}).get("name") or "Apartment search"

    # ---- Read me (filled last, but created first so it opens first)
    readme = B.wb.create_sheet("Read me")

    # ---- Weights
    wsh = B.sheet("Weights", "Weights", "Edit the yellow cells. Score and Rank on the Listings sheet recalculate; "
                  "only the ratios between weights matter.", ["Dimension", "Weight", "Share", "What it measures"],
                  ["text", "num1", "pct", "longtext"])
    wsh.column_dimensions["A"].width = 30
    wsh.column_dimensions["D"].width = 80
    w = crit.get("weights") or {}
    w0 = FIRST_ROW
    wN = w0 + len(DIMENSIONS) - 1
    for i, k in enumerate(DIMENSIONS):
        r = w0 + i
        B.put(wsh, r, 1, DIM_LABELS.get(k, k))
        c = B.put(wsh, r, 2, float(w.get(k, 0) or 0), "num1")
        c.fill = B.edit_fill
        B.put(wsh, r, 3, Formula(f"=IF(SUM($B${w0}:$B${wN})=0,0,B{r}/SUM($B${w0}:$B${wN}))"), "pct")
        B.put(wsh, r, 4, DIM_HELP.get(k, ""))
    r_neutral, r_target, r_rowhead, r_row = wN + 2, wN + 3, wN + 6, wN + 7
    B.put(wsh, r_neutral, 1, "Unknown dimension counts as")
    B.put(wsh, r_neutral, 2, 50.0, "num1").fill = B.edit_fill
    B.put(wsh, r_neutral, 4, "Neutral score for a dimension with no data (score.py uses 50); Range low/high on "
          "Listings show 0 and 100.")
    B.put(wsh, r_target, 1, "Target monthly cost")
    B.put(wsh, r_target, 2, float(_target(crit) or 0), "money").fill = B.edit_fill
    B.put(wsh, r_target, 4, "Drives the Vs target column (first budget tier in criteria.yaml).")
    preset = (crit.get("scoring") or {}).get("preset")
    B.put(wsh, r_target + 1, 1, "Weights from")
    B.put(wsh, r_target + 1, 4, f"criteria.yaml preset '{preset}'" if preset else "criteria.yaml weights")
    B.put(wsh, r_rowhead, 1, "Weights as a row (used by the formulas; do not edit)").font = B.Font(color=MUTED)
    for i, k in enumerate(DIMENSIONS):
        B.put(wsh, r_rowhead, 2 + i, DIM_LABELS.get(k, k)).font = B.Font(color=MUTED)
        B.put(wsh, r_row, 2 + i, Formula(f"=B{w0 + i}"), "num1").font = B.Font(color=MUTED)
    wrow = f"Weights!${_col(xl, 2)}${r_row}:${_col(xl, 1 + len(DIMENSIONS))}${r_row}"
    neutral = f"Weights!$B${r_neutral}"
    target = f"Weights!$B${r_target}"
    wsh.freeze_panes = "A5"
    counts["Weights"] = len(DIMENSIONS)

    # ---- Listings
    cols = listing_columns(crit)
    heads, kinds = [h for h, _, _ in cols], [k for _, k, _ in cols]
    sh = B.sheet("Listings", f"Listings: {search}",
                 "Dimension scores are values from score.py (blank = unknown). Score, Rank, Vs target and $ / sq ft "
                 "are live formulas driven by the Weights sheet. Ranks cover eligible listings only.", heads, kinds)
    idx = {h: i + 1 for i, h in enumerate(heads)}
    dim_cols = [i + 1 for i, k in enumerate(kinds) if k == "dim"]
    d0, d1 = _col(xl, dim_cols[0]), _col(xl, dim_cols[-1])
    n = len(active)
    last = FIRST_ROW + max(n, 1) - 1
    c_score, c_elig = _col(xl, idx["Score"]), _col(xl, idx["Eligible"])
    c_cost, c_sqft = _col(xl, idx["Monthly cost"]), _col(xl, idx["Sq ft"])
    for j, L in enumerate(active):
        r = FIRST_ROW + j
        for i, (_h, kind, get) in enumerate(cols, 1):
            if kind == "scoref":
                v: Any = Formula(
                    f"=IFERROR(IF(COUNT({d0}{r}:{d1}{r})=0,\"\",ROUND((SUMPRODUCT({d0}{r}:{d1}{r},{wrow})"
                    f"+{neutral}*SUMPRODUCT(({d0}{r}:{d1}{r}=\"\")*{wrow}))/SUM({wrow}),1)),\"\")")
            elif kind == "rankf":
                v = Formula(f"=IF(AND({c_elig}{r}=\"yes\",ISNUMBER({c_score}{r})),COUNTIFS(${c_elig}${FIRST_ROW}:"
                            f"${c_elig}${last},\"yes\",${c_score}${FIRST_ROW}:${c_score}${last},\">\"&{c_score}{r})+1,\"\")")
            elif kind == "vstargetf":
                v = Formula(f"=IF(AND(ISNUMBER({c_cost}{r}),{target}>0),{c_cost}{r}-{target},\"\")")
            elif kind == "psff":
                v = Formula(f"=IF(AND(ISNUMBER({c_cost}{r}),ISNUMBER({c_sqft}{r})),IF({c_sqft}{r}>0,"
                            f"{c_cost}{r}/{c_sqft}{r},\"\"),\"\")")
            else:
                v = get(L)
            B.put(sh, r, i, v, kind)
    B.finish(sh, len(cols), n, freeze=f"{_col(xl, idx['Unit'] + 1)}{FIRST_ROW}")
    green_scale = dict(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=50,
                       mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B")
    if n:
        sh.conditional_formatting.add(f"{c_score}{FIRST_ROW}:{c_score}{last}", ColorScaleRule(**green_scale))
        sh.conditional_formatting.add(f"{d0}{FIRST_ROW}:{d1}{last}", ColorScaleRule(**green_scale))
        ct = _col(xl, idx["Tier"])
        for name, color in _tier_colors(crit):
            lit = '"' + str(name).replace('"', '""') + '"'
            sh.conditional_formatting.add(f"{ct}{FIRST_ROW}:{ct}{last}", CellIsRule(
                operator="equal", formula=[lit], fill=B.Fill("solid", fgColor=color)))
        sh.conditional_formatting.add(f"{c_elig}{FIRST_ROW}:{c_elig}{last}", CellIsRule(
            operator="equal", formula=['"no"'], font=B.Font(color="9CA3AF")))
        cc = _col(xl, idx["Confidence"])
        sh.conditional_formatting.add(f"{cc}{FIRST_ROW}:{cc}{last}", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color="5B9BD5"))
        cv = _col(xl, idx["Vs target"])
        sh.conditional_formatting.add(f"{cv}{FIRST_ROW}:{cv}{last}", CellIsRule(
            operator="greaterThan", formula=["0"], font=B.Font(color="C00000")))
    counts["Listings"] = n

    # ---- Outreach: one row per building, questions from the missing facts
    focus = focus_listings(listings, top=top)
    groups: dict[tuple, list[dict]] = {}
    for L in focus:
        groups.setdefault(((L.get("building") or "").strip().lower(), (L.get("address") or "").strip().lower()),
                          []).append(L)
    o_heads = ["Building", "Address", "Units", "Best rank", "Questions to ask", "Contacted", "Channel",
               "Quoted gross", "Quoted net", "Lease term (mo)", "Fees confirmed", "Start date OK", "W/D confirmed",
               "Pets OK", "Tour date", "Status", "Next step"]
    o_kinds = ["text", "text", "text", "rankf", "longtext", "date", "text", "money", "money", "int", "text", "text",
               "text", "text", "date", "text", "longtext"]
    osh = B.sheet("Outreach", "Outreach: questions per building",
                  "Draft questions from the missing facts. The renter decides who to contact and sends every message.",
                  o_heads, o_kinds)
    osh.column_dimensions["E"].width = 70
    osh.column_dimensions["Q"].width = 30
    orow = FIRST_ROW
    for (_, _), Ls in sorted(groups.items(), key=lambda kv: min((L.get("score") or {}).get("rank") or 10 ** 6
                                                                 for L in kv[1])):
        L0 = Ls[0]
        qs: list[str] = []
        for L in Ls:
            for q in questions_for(L, crit):
                if q not in qs:
                    qs.append(q)
        ranks = [(L.get("score") or {}).get("rank") for L in Ls if (L.get("score") or {}).get("rank")]
        vals = [L0.get("building"), L0.get("address"), ", ".join(str(L.get("unit") or L["id"]) for L in Ls),
                min(ranks) if ranks else None, "\n".join(f"\u2022 {q}" for q in qs) or None]
        for i, v in enumerate(vals, 1):
            B.put(osh, orow, i, v, o_kinds[i - 1])
        for i in range(6, len(o_heads) + 1):
            B.put(osh, orow, i, None, o_kinds[i - 1]).fill = B.edit_fill
        orow += 1
    o_n = orow - FIRST_ROW
    o_last = max(orow - 1, FIRST_ROW)
    for col, choices in (("G", "email,phone,portal,in person,text"), ("K", "yes,no,partly"),
                         ("L", "yes,no"), ("M", "yes,no"), ("N", "yes,no,n/a"),
                         ("P", "to contact,contacted,replied,tour booked,toured,applied,passed")):
        dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
        osh.add_data_validation(dv)
        dv.add(f"{col}{FIRST_ROW}:{col}{o_last + 50}")
    B.finish(osh, len(o_heads), o_n, freeze=f"C{FIRST_ROW}")
    counts["Outreach"] = o_n

    # ---- To verify: active listings with missing facts, best first
    v_heads = ["Rank", "Building", "Unit", "Missing facts", "Score low", "Score high", "Confidence", "Checked",
               "Days since check", "Listing", "ID"]
    v_kinds = ["rankf", "text", "text", "longtext", "num1", "num1", "pct", "date", "int", "url", "text"]
    vsh = B.sheet("To verify", "To verify: missing facts",
                  "Unknown values stay blank and are never estimated. Score low/high show where each listing could land "
                  "once the facts are in.", v_heads, v_kinds)
    vrow = FIRST_ROW
    for L in active:
        miss = missing_facts(L, crit)
        if not miss:
            continue
        s = L.get("score") or {}
        rng = s.get("range") or [None, None]
        chk = _date_or_text(L.get("checked"))
        vals = [s.get("rank"), L.get("building"), L.get("unit"), ", ".join(miss), rng[0], rng[1],
                (s.get("confidence") or 0) / 100, chk,
                Formula(f"=IF(ISNUMBER(H{vrow}),TODAY()-H{vrow},\"\")") if isinstance(chk, dt.date) else None,
                _url(L), L["id"]]
        for i, v in enumerate(vals, 1):
            B.put(vsh, vrow, i, v, v_kinds[i - 1])
        vrow += 1
    v_n = vrow - FIRST_ROW
    if v_n:
        vsh.conditional_formatting.add(f"I{FIRST_ROW}:I{vrow - 1}", CellIsRule(
            operator="greaterThan", formula=["14"], font=B.Font(color="C00000")))
    B.finish(vsh, len(v_heads), v_n, freeze=f"D{FIRST_ROW}")
    counts["To verify"] = v_n

    # ---- Upfront charges: what it takes to sign, known items only
    u_heads = ["Rank", "Building", "Unit", "First month", "Security deposit", "Broker fee", "One-time fees",
               "Application fee", "Move-in fee", "Known upfront (at least)", "Still unknown", "Free (mo)",
               "Lease (mo)", "Listing", "ID"]
    u_kinds = ["rankf", "text", "text", "money", "money", "money", "money", "money", "money", "money", "text", "num1",
               "int", "url", "text"]
    ush = B.sheet("Upfront charges", "Upfront charges at signing",
                  "Known items only; blank means unknown, not zero. Ask the leasing office for anything in 'Still "
                  "unknown'.", u_heads, u_kinds)
    ush.column_dimensions["K"].width = 34
    urow = FIRST_ROW
    for L in active:
        if not (L.get("score") or {}).get("eligible") and L.get("status") not in ("shortlist", "touring", "applied"):
            continue
        p = L.get("price") or {}
        dep = p.get("deposit", p.get("security_deposit"))
        items = {"First month": p.get("gross"), "Security deposit": dep, "Broker fee": p.get("broker_fee"),
                 "One-time fees": p.get("fees_one_time"), "Application fee": p.get("application_fee"),
                 "Move-in fee": p.get("move_in_fee")}
        unknown = [k.lower() for k, v in items.items() if v is None]
        vals = [(L.get("score") or {}).get("rank"), L.get("building"), L.get("unit"), *items.values(),
                Formula(f"=IF(COUNT(D{urow}:I{urow})=0,\"\",SUM(D{urow}:I{urow}))"), ", ".join(unknown) or None,
                p.get("free_months"), p.get("lease_months"), _url(L), L["id"]]
        for i, v in enumerate(vals, 1):
            B.put(ush, urow, i, v, u_kinds[i - 1])
        urow += 1
    B.finish(ush, len(u_heads), urow - FIRST_ROW, freeze=f"D{FIRST_ROW}")
    counts["Upfront charges"] = urow - FIRST_ROW

    # ---- Rejected or held
    r_heads = ["Building", "Unit", "Status", "Why", "Monthly cost", "Score", "Notes", "Checked", "Listing", "ID"]
    r_kinds = ["text", "text", "text", "longtext", "money", "num1", "longtext", "date", "url", "text"]
    rsh = B.sheet("Rejected or held", "Rejected or held",
                  "Rejected by the renter, or held back by a gate (over the hard budget cap or a failed must-have). "
                  "Kept so nothing is silently dropped.", r_heads, r_kinds)
    rrow = FIRST_ROW
    hard = (crit.get("budget") or {}).get("hard_max")
    for L in rows:
        s = L.get("score") or {}
        held = is_active(L) and not s.get("eligible")
        if L.get("status") != "rejected" and not held:
            continue
        why = []
        for g in s.get("gates") or []:
            if g == "over_budget":
                why.append(f"over the hard cap (${s.get('monthly_cost') or 0:,.0f} > ${hard or 0:,.0f})")
            elif g.startswith("fails:"):
                why.append("fails must-have: " + g.split(":", 1)[1].replace(",", ", "))
        if L.get("status") == "rejected":
            why.insert(0, "rejected")
        vals = [L.get("building"), L.get("unit"), L.get("status") if L.get("status") == "rejected" else "held",
                "; ".join(why) or None, s.get("monthly_cost"), s.get("overall"), L.get("notes"),
                _date_or_text(L.get("checked")), _url(L), L["id"]]
        for i, v in enumerate(vals, 1):
            B.put(rsh, rrow, i, v, r_kinds[i - 1])
        rrow += 1
    B.finish(rsh, len(r_heads), rrow - FIRST_ROW, freeze=f"C{FIRST_ROW}")
    counts["Rejected or held"] = rrow - FIRST_ROW

    # ---- Gone fast (only with history)
    if history:
        g_heads = ["Building", "Unit", "Listed", "Gone", "Days listed", "Last gross", "Beds", "Sq ft",
                   "Neighborhood", "Listing", "ID"]
        g_kinds = ["text", "text", "date", "date", "int", "money", "num1", "sqft", "text", "url", "text"]
        gsh = B.sheet("Gone fast", "Gone fast", "Units that left the market, quickest first: a guide to how fast "
                      "to decide. Days are between the first and last sighting, not a contract date.", g_heads, g_kinds)
        grow = FIRST_ROW
        for g in history:
            L = by_id.get(g["id"]) or {"id": g["id"]}
            vals = [L.get("building"), L.get("unit"), _date_or_text(g.get("listed")), _date_or_text(g.get("gone")),
                    g.get("days"), g.get("last_price"), L.get("beds"), L.get("sqft"), L.get("neighborhood"), _url(L),
                    g["id"]]
            for i, v in enumerate(vals, 1):
                B.put(gsh, grow, i, v, g_kinds[i - 1])
            grow += 1
        B.finish(gsh, len(g_heads), grow - FIRST_ROW, freeze=f"C{FIRST_ROW}")
        counts["Gone fast"] = grow - FIRST_ROW

    # ---- Sources
    s_heads = ["Building", "Unit", "Kind", "Label", "URL", "Checked", "ID"]
    s_kinds = ["text", "text", "text", "text", "url", "date", "text"]
    ssh = B.sheet("Sources", "Sources", "Where each fact came from. Reviews and 311 reports are allegations, not "
                  "findings.", s_heads, s_kinds)
    srow = FIRST_ROW
    for L in rows:
        for kind, label, url in _sources(L):
            vals = [L.get("building"), L.get("unit"), kind, label, url, _date_or_text(L.get("checked")), L["id"]]
            for i, v in enumerate(vals, 1):
                B.put(ssh, srow, i, v, s_kinds[i - 1])
            srow += 1
    B.finish(ssh, len(s_heads), srow - FIRST_ROW, freeze=f"C{FIRST_ROW}")
    counts["Sources"] = srow - FIRST_ROW

    # ---- Read me
    checked = sorted(d for d in (parse_date(L.get("checked")) for L in active) if d)
    lines = [
        (search, None),
        (f"Exported {today.isoformat()} by ApartmentScout from listings.json and criteria.yaml.", None),
        ("", None),
        ("Listings", f"{counts['Listings']} active ({sum(1 for L in active if (L.get('score') or {}).get('rank'))} "
                     f"ranked); {len(listings)} in the search in total"),
        ("Data checked", f"{checked[0].isoformat()} to {checked[-1].isoformat()}" if checked else "no check dates"),
        ("", None),
        ("How to use", "Change the yellow cells on Weights; Score and Rank on Listings recalculate."),
        ("Monthly cost", "Net-effective rent (gross spread over the lease after free months) + monthly fees + "
                         "one-time fees spread over the lease."),
        ("Unknowns", "Blank means unknown and is never estimated. An unknown dimension counts as the neutral value on "
                     "Weights; Range low/high show the worst and best case."),
        ("Gates", "Ranks cover eligible listings only: within the hard budget cap, no failed must-have, not gone or "
                  "rejected. Held listings are on 'Rejected or held'."),
        ("Honesty", "Observations are not lease offers. Confirm price, fees and dates in writing before relying on "
                    "them. Reviews and 311 reports are allegations."),
        ("", None),
        *[(title, f"{counts[title]} rows") for title in counts if title != "Weights"],
    ]
    readme["A1"].font = B.Font(bold=True, size=14)
    for i, (a, b) in enumerate(lines, 1):
        B.put(readme, i, 1, a)
        if b is not None:
            B.put(readme, i, 2, b)
            readme.cell(i, 1).font = B.Font(bold=True)
    readme["A1"].font = B.Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 18
    readme.column_dimensions["B"].width = 110

    B.wb.calculation.fullCalcOnLoad = True   # Excel and LibreOffice compute every formula on open
    B.wb.active = B.wb.sheetnames.index("Listings")
    path.parent.mkdir(parents=True, exist_ok=True)
    B.wb.save(path)
    return counts


def _sources(L: dict) -> list[tuple[str, str | None, str]]:
    out: list[tuple[str, str | None, str]] = []
    seen: set[str] = set()

    def add(kind: str, label: Any, url: Any) -> None:
        if isinstance(url, str) and url.startswith(URL_SCHEMES) and url not in seen:
            seen.add(url)
            out.append((kind, str(label) if label else None, url))

    for k, u in (L.get("links") or {}).items():
        add(k, None, u)
    for s in L.get("sources") or []:
        if isinstance(s, dict):
            add("source", s.get("label") or s.get("name"), s.get("url"))
        else:
            add("source", None, s)
    rv = L.get("reviews") or {}
    for r in rv.get("ratings") or []:
        if isinstance(r, dict):
            add("review", " ".join(str(x) for x in (r.get("label"), r.get("value")) if x), r.get("url"))
    for s in rv.get("sources") or []:
        if isinstance(s, dict):
            add("review source", s.get("label"), s.get("url"))
    return out
